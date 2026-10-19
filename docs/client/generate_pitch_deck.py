"""MarketLoopAI — EAII second-round pitch deck (7 slides, 16:9)."""

from __future__ import annotations

import math
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

DIR = Path(__file__).resolve().parent
OUT = DIR / "MarketLoopAI-EAII-Pitch-Deck.pptx"
LOGO = DIR / "marketloopai-logo.png"

BLACK = RGBColor(0x0A, 0x0A, 0x0A)
NEAR_BLACK = RGBColor(0x14, 0x14, 0x16)
CARD = RGBColor(0x1A, 0x1A, 0x1E)
RED = RGBColor(0xD4, 0x06, 0x06)
RED_HOT = RGBColor(0xFA, 0x07, 0x07)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SILVER = RGBColor(0xC7, 0xC9, 0xCC)
DIM = RGBColor(0x6E, 0x6E, 0x73)

W = Inches(13.333)
H = Inches(7.5)

STORY = ["Cover", "The loop", "Market", "Traction", "GTM", "Team", "The ask"]


def render_logo(dest: Path, size: int = 1024) -> None:
    """Brand mark: outer red loop + inner silver loop (from site SVG)."""
    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    s = size / 72.0

    def px(x: float, y: float):
        return (x * s, y * s)

    def ang(x: float, y: float, cx: float = 36.0, cy: float = 36.0) -> float:
        return math.degrees(math.atan2(y - cy, x - cx))

    d.arc(
        [12 * s, 12 * s, 60 * s, 60 * s],
        start=ang(40.17, 59.64),
        end=ang(59.64, 40.17),
        fill=(250, 7, 7, 255),
        width=max(3, int(7 * s)),
    )
    d.polygon([px(64.56, 41.04), px(61.44, 34.22), px(56.2, 43.35)], fill=(250, 7, 7, 255))
    d.arc(
        [23 * s, 23 * s, 49 * s, 49 * s],
        start=ang(33.74, 23.2),
        end=ang(23.2, 33.74),
        fill=(199, 201, 204, 255),
        width=max(2, int(5.5 * s)),
    )
    d.polygon([px(18.77, 32.96), px(23.21, 30.31), px(25, 36)], fill=(199, 201, 204, 255))
    r = 3.4 * s
    cx, cy = 36 * s, 36 * s
    d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=(255, 255, 255, 255))
    dest.parent.mkdir(parents=True, exist_ok=True)
    img.save(dest)


def set_run(run, size, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs", "a:ea"):
        node = rPr.find(qn(tag))
        if node is None and tag == "a:latin":
            node = etree.SubElement(rPr, qn(tag))
        if node is not None:
            node.set("typeface", font)


def add_text(
    slide,
    l,
    t,
    w,
    h,
    text,
    size=14,
    bold=False,
    color=WHITE,
    align=PP_ALIGN.LEFT,
    font="Calibri",
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color, font)
    return box


def bullets(slide, l, t, w, h, lines, size=13, color=SILVER, gap=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = "▸  " + line
        set_run(run, size, False, color)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_shape(s, color)
    return s


def rrect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill_shape(s, color)
    try:
        s.adjustments[0] = 0.08
    except Exception:
        pass
    return s


def paint_bg(slide, color=BLACK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_logo(slide, l, t, size):
    slide.shapes.add_picture(str(LOGO), l, t, size, size)


def story_bar(slide, current: int):
    """Narrative tracker so the deck reads as one flow, not seven posters."""
    x0 = Inches(0.5)
    y = Inches(0.18)
    gap = Inches(1.82)
    for i, label in enumerate(STORY):
        x = x0 + gap * i
        active = i == current
        color = RED if active else DIM
        add_text(slide, x, y, Inches(1.7), Inches(0.22), label.upper(), size=9, bold=active, color=color)
        if i < len(STORY) - 1:
            add_text(slide, x + Inches(1.55), y, Inches(0.28), Inches(0.22), "→", size=9, color=RED if i < current else DIM)


def chrome(slide, page: int, total: int = 7):
    rect(slide, Inches(0), Inches(0), Inches(0.08), H, RED)
    if page > 1:
        story_bar(slide, page - 1)
        add_logo(slide, Inches(12.55), Inches(0.12), Inches(0.42))
    add_text(
        slide,
        Inches(0.45),
        Inches(7.18),
        Inches(9.5),
        Inches(0.25),
        "MarketLoopAI  ·  Confidential  ·  EAII Second-Round Screening",
        size=10,
        color=DIM,
    )
    add_text(slide, Inches(11.6), Inches(7.18), Inches(1.4), Inches(0.25), f"{page}  /  {total}", size=10, color=DIM, align=PP_ALIGN.RIGHT)


def kicker(slide, text, t=Inches(0.42)):
    add_text(slide, Inches(0.5), t, Inches(12), Inches(0.28), text.upper(), size=11, bold=True, color=RED)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def style_table(table, header=True):
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.06)
            cell.margin_bottom = Inches(0.06)
            fill = BLACK if (header and i == 0) else (CARD if i % 2 else NEAR_BLACK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(tag))
                ln.set("w", "0")
                etree.SubElement(ln, qn("a:noFill"))
            for p in cell.text_frame.paragraphs:
                p.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(12 if i else 11)
                    run.font.bold = i == 0 or cell is row.cells[0]
                    run.font.color.rgb = WHITE if i == 0 else (SILVER if cell is not row.cells[0] else WHITE)


def build():
    render_logo(LOGO)
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # ─── 1 COVER ───────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    rect(s, Inches(0), Inches(0), Inches(0.12), H, RED)
    rect(s, Inches(0), Inches(0), W, Inches(0.06), RED)
    add_logo(s, Inches(5.55), Inches(1.15), Inches(2.2))
    add_text(s, Inches(0.5), Inches(3.45), Inches(12.3), Inches(0.45),
             "MARKETLOOPAI", size=16, bold=True, color=RED, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.8), Inches(3.95), Inches(11.7), Inches(1.15),
             "One idea in.\nCreate, publish, measure, improve.",
             size=34, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(1.5), Inches(5.2), Inches(10.3), Inches(0.4),
             "Content, video, campaigns, and performance — one loop across every platform.",
             size=16, color=SILVER, align=PP_ALIGN.CENTER)

    rrect(s, Inches(2.4), Inches(5.8), Inches(8.5), Inches(0.85), CARD)
    add_text(s, Inches(2.55), Inches(5.9), Inches(8.2), Inches(0.32),
             "Fasiledes Shimelis   ·   Founder & CEO          Hlina Sisay   ·   Co-founder",
             size=14, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2.55), Inches(6.22), Inches(8.2), Inches(0.32),
             "EAII AI Startup Support Program   ·   Second-round pitch   ·   September 2026",
             size=12, color=DIM, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.28),
             "marketloopai.com", size=12, bold=True, color=RED, align=PP_ALIGN.CENTER)
    notes(s,
          "Pause on the cover two seconds. Say the name: MarketLoopAI. Then the line: one idea in — create, publish, measure, improve. "
          "Introduce yourselves: Fasiledes Shimelis, Founder and CEO, and Hlina Sisay, Co-founder. Then: we are applying to EAII to take this from a live waitlist product to paid video generation.")

    # ─── 2 PROBLEM → SOLUTION (the loop) ───────────────────────
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    chrome(s, 2)
    kicker(s, "Problem  →  Solution")
    add_text(s, Inches(0.5), Inches(0.68), Inches(12.3), Inches(0.42),
             "Marketers do not need another chatbot. They need create → publish → learn.",
             size=20, bold=True)

    pains = [
        ("Fragmented", "Writer, editor, scheduler, five ad managers. The first draft still ships."),
        ("Blind spend", "Campaigns go live. Performance sits in five dashboards. Nobody feeds winners back into the next ad."),
        ("Video is stuck", "The ad unit of 2026 still costs thousands and weeks. Most SMBs skip it."),
    ]
    for i, (title, body) in enumerate(pains):
        x = Inches(0.5 + i * 4.2)
        rrect(s, x, Inches(1.18), Inches(4.0), Inches(1.38), CARD)
        rect(s, x, Inches(1.18), Inches(0.07), Inches(1.38), RED)
        add_text(s, x + Inches(0.22), Inches(1.26), Inches(3.6), Inches(0.3), title, size=14, bold=True)
        add_text(s, x + Inches(0.22), Inches(1.56), Inches(3.6), Inches(0.88), body, size=12, color=SILVER)

    add_text(s, Inches(0.5), Inches(2.68), Inches(12.3), Inches(0.24),
             "CREATE", size=10, bold=True, color=RED)
    create = [
        ("Idea", "Prompt or product"),
        ("Generate", "Up to 9 versions"),
        ("Rank", "Humanize + score"),
        ("Adapt", "Each platform’s voice"),
        ("Video", "Clip or storyboard"),
    ]
    for i, (t, d) in enumerate(create):
        x = Inches(0.5 + i * 2.52)
        rrect(s, x, Inches(2.94), Inches(2.38), Inches(1.22), CARD)
        add_text(s, x + Inches(0.08), Inches(3.02), Inches(2.22), Inches(0.28), t, size=13, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.08), Inches(3.32), Inches(2.22), Inches(0.7), d, size=11, color=SILVER, align=PP_ALIGN.CENTER)
        if i < len(create) - 1:
            add_text(s, x + Inches(2.22), Inches(3.28), Inches(0.32), Inches(0.28), "→", size=14, color=RED, align=PP_ALIGN.CENTER)

    add_text(s, Inches(0.5), Inches(4.28), Inches(12.3), Inches(0.24),
             "OPERATE  ·  THE DATA COMES BACK IN", size=10, bold=True, color=RED)
    operate = [
        ("Publish", "LinkedIn, TikTok, Snapchat, X, Facebook, Instagram"),
        ("Run ads", "TikTok, Snapchat, X, Google, Meta Ads — one workspace"),
        ("Analyze", "Unified performance across content and campaigns"),
        ("Suggest", "Next creative from what actually won"),
    ]
    for i, (t, d) in enumerate(operate):
        x = Inches(0.5 + i * 3.15)
        rrect(s, x, Inches(4.54), Inches(3.0), Inches(1.18), CARD)
        add_text(s, x + Inches(0.12), Inches(4.62), Inches(2.76), Inches(0.28), t, size=13, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.12), Inches(4.92), Inches(2.76), Inches(0.68), d, size=11, color=SILVER, align=PP_ALIGN.CENTER)
        if i < len(operate) - 1:
            add_text(s, x + Inches(2.85), Inches(4.88), Inches(0.32), Inches(0.28), "→", size=14, color=RED, align=PP_ALIGN.CENTER)

    rrect(s, Inches(0.5), Inches(5.88), Inches(12.3), Inches(0.88), CARD)
    add_text(s, Inches(0.7), Inches(6.02), Inches(12.0), Inches(0.62),
             "That is the brand: content flows out to every channel. Performance flows back in. Winners become the next brief — not another blank prompt.",
             size=14, color=SILVER)
    notes(s,
          "Walk CREATE then OPERATE. Generation is only half the product. We publish to LinkedIn, TikTok, Snapchat, X, Facebook, Instagram. We run campaigns on TikTok Ads, Snapchat Ads, X Ads, Google Ads, and Meta Ads. Then we pull performance into one view and suggest the next creative from what won. That closed loop is the logo.")

    # ─── 3 MARKET ──────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    chrome(s, 3)
    kicker(s, "Market opportunity")
    add_text(s, Inches(0.5), Inches(0.68), Inches(12.3), Inches(0.5),
             "Short-form video is the ad. SMBs still cannot produce it.",
             size=24, bold=True)

    metrics = [
        ("$847M", "AI video generator market, 2026"),
        ("$3.35B", "Projected by 2034  ·  18.8% CAGR"),
        ("46%", "Text-to-video share of that market"),
    ]
    for i, (v, l) in enumerate(metrics):
        x = Inches(0.5 + i * 4.2)
        rrect(s, x, Inches(1.35), Inches(4.0), Inches(1.55), CARD)
        add_text(s, x + Inches(0.25), Inches(1.48), Inches(3.5), Inches(0.6), v, size=26, bold=True, color=RED)
        add_text(s, x + Inches(0.25), Inches(2.1), Inches(3.5), Inches(0.55), l, size=13, color=SILVER)

    add_text(s, Inches(0.5), Inches(3.05), Inches(12), Inches(0.28),
             "Source: Fortune Business Insights, AI Video Generator Market (2026–2034).",
             size=11, color=DIM)

    wedges = [
        ("TAM", "AI video generation + AI marketing tools businesses already buy as SaaS."),
        ("SAM", "SMBs, e-commerce, SaaS, agencies who need posts, short video ads, and campaign ops in one place."),
        ("SOM", "Waitlist → paid beta in 12 months. Founder-led, then agencies as the multiplier."),
    ]
    for i, (t, d) in enumerate(wedges):
        y = Inches(3.45 + i * 0.78)
        rrect(s, Inches(0.5), y, Inches(6.3), Inches(0.7), CARD)
        add_text(s, Inches(0.7), y + Inches(0.16), Inches(0.9), Inches(0.4), t, size=14, bold=True, color=RED)
        add_text(s, Inches(1.7), y + Inches(0.12), Inches(4.9), Inches(0.5), d, size=12, color=SILVER)

    rrect(s, Inches(7.1), Inches(3.45), Inches(5.7), Inches(3.15), CARD)
    add_text(s, Inches(7.35), Inches(3.6), Inches(5.2), Inches(0.32), "Why now", size=16, bold=True)
    bullets(
        s,
        Inches(7.35),
        Inches(4.05),
        Inches(5.2),
        Inches(2.3),
        [
            "Wan 2.2 (Apache 2.0) made commercial self-hosting possible without per-clip lock-in.",
            "Algorithms reward native video. Static posts lose reach.",
            "SMBs already pay $29–$199/mo for point tools. They will pay for a loop that ships and learns.",
            "Agencies need one workspace to publish, run ads, and see what won — not six logins.",
        ],
        size=13,
        gap=8,
    )
    notes(s,
          "Cite Fortune: $847M in 2026 to $3.35B by 2034. We are not claiming the whole market. Beachhead: SMBs and agencies who already buy tools but cannot ship video weekly. Why now: open-weight video models plus short video as the ad unit.")

    # ─── 4 TRACTION + BUSINESS MODEL ───────────────────────────
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    chrome(s, 4)
    kicker(s, "Traction  ·  Business model")
    add_text(s, Inches(0.5), Inches(0.68), Inches(12.3), Inches(0.45),
             "Pre-revenue. Product surface is live. Engine is specified.",
             size=22, bold=True)

    facts = [
        ("2026", "Founded. Building in public."),
        ("Live", "marketloopai.com — waitlist, CMS, legal, security."),
        ("6 + 5", "Publish on 6 socials. Run ads on 5 networks."),
        ("The loop", "Create → publish → analyze → suggest the next winner."),
    ]
    for i, (v, l) in enumerate(facts):
        x = Inches(0.5 + i * 3.15)
        rrect(s, x, Inches(1.28), Inches(3.0), Inches(1.4), CARD)
        add_text(s, x + Inches(0.18), Inches(1.38), Inches(2.65), Inches(0.4), v, size=18, bold=True, color=RED)
        add_text(s, x + Inches(0.18), Inches(1.82), Inches(2.65), Inches(0.7), l, size=12, color=SILVER)

    add_text(s, Inches(0.5), Inches(2.85), Inches(6), Inches(0.32), "SaaS plans (published)", size=15, bold=True)
    rows = [
        ["Plan", "Price / mo", "AI videos*", "Who it is for"],
        ["Free", "$0", "1", "Try the loop"],
        ["Starter", "$29", "5", "Solo creators"],
        ["Pro  ·  featured", "$79", "20", "Teams + campaigns"],
        ["Agency", "$199", "80", "Clients + analytics"],
    ]
    table_shape = s.shapes.add_table(len(rows), 4, Inches(0.5), Inches(3.22), Inches(7.4), Inches(3.35)).table
    table_shape.columns[0].width = Inches(2.0)
    table_shape.columns[1].width = Inches(1.5)
    table_shape.columns[2].width = Inches(1.5)
    table_shape.columns[3].width = Inches(2.4)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table_shape.cell(r, c).text = val
    style_table(table_shape)

    rrect(s, Inches(8.15), Inches(2.85), Inches(4.7), Inches(3.72), CARD)
    add_text(s, Inches(8.4), Inches(3.0), Inches(4.25), Inches(0.32), "How we make money", size=15, bold=True)
    bullets(
        s,
        Inches(8.4),
        Inches(3.4),
        Inches(4.25),
        Inches(2.95),
        [
            "Subscriptions first. Credits under the hood so video can cost more than a caption.",
            "Failed GPU jobs refund credits.",
            "Later: agency seats, extra credit packs, premium video tiers.",
            "Pro / Agency: campaign management, analytics, winner-based suggestions.",
            "No paid ads until GPU-seconds prove contribution margin.",
            "*Site video counts map to credits as models get costlier.",
        ],
        size=12,
        gap=7,
    )
    notes(s,
          "Honest traction: live branded product, waitlist plus referral, complete video architecture — not fake ARR. "
          "Pro at $79 is the hero plan. Agency is expansion. Credits protect GPU cost.")

    # ─── 5 GTM + COMPETITION ───────────────────────────────────
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    chrome(s, 5)
    kicker(s, "Go-to-market  ·  Competition")
    add_text(s, Inches(0.5), Inches(0.68), Inches(12.3), Inches(0.4),
             "Design partners first. Agencies as the multiplier.",
             size=22, bold=True)

    gtm = [
        ("1  Waitlist", "Referral codes. Show the full loop on TikTok and LinkedIn — not the model."),
        ("2  Partners", "10–20 ecom, SaaS, and agency teams. Paid beta on create + campaigns."),
        ("3  Product-led", "Free → Starter when credits run out. Pro unlocks campaign management."),
        ("4  Agencies", "One workspace per client: publish, ads, analytics. Highest LTV."),
    ]
    for i, (t, d) in enumerate(gtm):
        x = Inches(0.5 + i * 3.15)
        rrect(s, x, Inches(1.2), Inches(3.0), Inches(1.65), CARD)
        add_text(s, x + Inches(0.18), Inches(1.32), Inches(2.65), Inches(0.35), t, size=13, bold=True, color=RED)
        add_text(s, x + Inches(0.18), Inches(1.7), Inches(2.65), Inches(0.95), d, size=12, color=SILVER)

    add_text(s, Inches(0.5), Inches(3.02), Inches(12), Inches(0.28),
             "Not another chatbot. Not a scheduler with AI bolted on. The loop includes campaigns and performance.", size=13, color=SILVER)

    rows = [
        ["", "Jasper / Copy.ai", "Runway / Pika", "Schedulers", "MarketLoopAI"],
        ["Ranked creative", "No", "No", "No", "Yes"],
        ["Publish to social", "No", "No", "Yes", "Yes · 6 networks"],
        ["Ad campaign management", "No", "No", "Partial", "Yes · 5 ad platforms"],
        ["Unified analytics", "No", "No", "Partial", "Yes"],
        ["Suggest from winners", "No", "No", "No", "Yes"],
        ["AI video in-product", "No", "Yes", "No", "Yes · self-hosted"],
    ]
    table_shape = s.shapes.add_table(len(rows), 5, Inches(0.5), Inches(3.35), Inches(12.35), Inches(3.4)).table
    for i, w in enumerate([2.15, 2.4, 2.4, 2.4, 3.0]):
        table_shape.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table_shape.cell(r, c).text = val
    style_table(table_shape)
    notes(s,
          "GTM in four beats: waitlist, design partners, PLG, agencies. Pro and Agency unlock campaign management. "
          "Competition: Jasper owns copy, Runway owns experimental video, schedulers own calendars. Nobody owns create, publish, analyze, and suggest-from-winners in one loop.")

    # ─── 6 TEAM ────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    chrome(s, 6)
    kicker(s, "Core team and advisors")
    add_text(s, Inches(0.5), Inches(0.68), Inches(12.3), Inches(0.4),
             "Two founders. Complementing skills. Advisors next.",
             size=22, bold=True)

    # CEO
    rrect(s, Inches(0.5), Inches(1.25), Inches(6.05), Inches(3.55), CARD)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(1.5), Inches(1.05), Inches(1.05))
    fill_shape(circ, RED)
    add_text(s, Inches(0.8), Inches(1.72), Inches(1.05), Inches(0.6), "FS", size=20, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(2.05), Inches(1.55), Inches(4.2), Inches(0.4), "Fasiledes Shimelis", size=20, bold=True)
    add_text(s, Inches(2.05), Inches(1.98), Inches(4.2), Inches(0.32), "Founder  ·  CEO", size=14, color=RED)
    add_text(s, Inches(0.8), Inches(2.75), Inches(5.45), Inches(1.8),
             "Owns product, engineering, and company vision. Shipping the MarketLoopAI web product and the self-hosted Wan 2.2 video architecture.\n\n"
             "Available ≥ 3 days/week for the full 4–6 month EAII program.",
             size=13, color=SILVER)

    # Co-founder
    rrect(s, Inches(6.8), Inches(1.25), Inches(6.05), Inches(3.55), CARD)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.1), Inches(1.5), Inches(1.05), Inches(1.05))
    fill_shape(circ, RGBColor(0xC7, 0xC9, 0xCC))
    add_text(s, Inches(7.1), Inches(1.72), Inches(1.05), Inches(0.6), "HS", size=20, bold=True, color=BLACK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(8.35), Inches(1.55), Inches(4.2), Inches(0.4), "Hlina Sisay", size=20, bold=True)
    add_text(s, Inches(8.35), Inches(1.98), Inches(4.2), Inches(0.32), "Co-founder", size=14, color=RED)
    add_text(s, Inches(7.1), Inches(2.75), Inches(5.45), Inches(1.8),
             "Owns go-to-market, partnerships, and customer conversations. Turns the waitlist into design partners. Owns how campaigns get published, read, and turned into the next brief.\n\n"
             "Available ≥ 3 days/week for the full 4–6 month EAII program.",
             size=13, color=SILVER)

    rrect(s, Inches(0.5), Inches(5.0), Inches(6.05), Inches(1.75), CARD)
    add_text(s, Inches(0.75), Inches(5.15), Inches(5.55), Inches(0.3), "Advisors we will lock in-program", size=14, bold=True, color=RED)
    add_text(s, Inches(0.75), Inches(5.5), Inches(5.55), Inches(1.05),
             "Growth / performance marketing  ·  AI infrastructure (GPU ops and inference cost). Unpaid. Recruited during EAII.",
             size=13, color=SILVER)

    rrect(s, Inches(6.8), Inches(5.0), Inches(6.05), Inches(1.75), CARD)
    add_text(s, Inches(7.05), Inches(5.15), Inches(5.55), Inches(0.3), "First hires after beta", size=14, bold=True, color=RED)
    add_text(s, Inches(7.05), Inches(5.5), Inches(5.55), Inches(1.05),
             "Customer success for design partners  ·  Full-stack engineer for worker reliability. Mentorship bridges until those seats are funded.",
             size=13, color=SILVER)
    notes(s,
          "Introduce Fasiledes Shimelis, Founder and CEO — product and engine. Then Hlina Sisay, Co-founder — GTM and customers. Both commit three-plus days a week. Advisors and first hires come through the program. Do not oversell a board you do not have.")

    # ─── 7 FINANCIALS + ASK ────────────────────────────────────
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    chrome(s, 7)
    kicker(s, "Financial projection  ·  The ask")
    add_text(s, Inches(0.5), Inches(0.68), Inches(12.3), Inches(0.4),
             "Conservative ramp. Compute before paid acquisition.",
             size=22, bold=True)

    rows = [
        ["", "2026  (now)", "2027", "2028"],
        ["Stage", "Pre-revenue  ·  beta", "Paid launch", "Scale"],
        ["Paying accounts (YE)", "0–20 design partners", "~180", "~1,200"],
        ["Blended ARPU / mo", "—", "$55", "$62"],
        ["ARR (YE, projected)", "$0", "~$120k", "~$0.9M"],
        ["Gross margin target", "n/a until GPU measured", "55%+", "65%+"],
    ]
    table_shape = s.shapes.add_table(len(rows), 4, Inches(0.5), Inches(1.2), Inches(7.55), Inches(3.25)).table
    table_shape.columns[0].width = Inches(2.15)
    for i in range(1, 4):
        table_shape.columns[i].width = Inches(1.8)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table_shape.cell(r, c).text = val
    style_table(table_shape)
    add_text(s, Inches(0.5), Inches(4.55), Inches(7.55), Inches(0.55),
             "Projections, not bookings. Credit prices freeze only after GPU-seconds are measured. We will not buy ads until contribution margin is positive.",
             size=12, color=DIM)

    rrect(s, Inches(8.25), Inches(1.2), Inches(4.55), Inches(5.35), CARD)
    add_text(s, Inches(8.5), Inches(1.38), Inches(4.1), Inches(0.28), "THE ASK", size=13, bold=True, color=RED)
    add_text(s, Inches(8.5), Inches(1.72), Inches(4.1), Inches(0.95),
             "Acceptance into the EAII AI Startup Support Program (4–6 months).",
             size=16, bold=True)
    bullets(
        s,
        Inches(8.5),
        Inches(2.75),
        Inches(4.1),
        Inches(2.0),
        [
            "GPU / cloud compute to run Wan 2.2 in production",
            "GTM mentorship to convert waitlist → paid partners",
            "Advisor and customer introductions",
            "Help hiring the first customer-success seat",
        ],
        size=13,
        gap=6,
    )
    add_text(s, Inches(8.5), Inches(4.9), Inches(4.1), Inches(1.15),
             "Time split  ·  40% product  ·  30% GTM  ·  20% infra  ·  10% ops\n\n"
             "Fasiledes + Hlina: ≥ 3 days/week for the full program.\n"
             "marketloopai.com",
             size=12, color=SILVER)
    notes(s,
          "Close: ~180 paying accounts in 2027 at about $55 blended ARPU is about $120k ARR. 2028 is agency scale. "
          "The ask is the program plus GPU compute. Both founders, three days a week, four to six months. "
          "Last line: MarketLoopAI turns one idea into ranked content, published campaigns, and the next ad based on what actually won.")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
